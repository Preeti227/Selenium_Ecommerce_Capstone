"""Read and validate the same test fields from JSON or Excel."""
import json
import os
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
FIELDS = ("base_url", "product", "quantity", "first_name", "last_name", "telephone")


def load_data():
    path = Path(os.getenv("TEST_DATA", str(ROOT / "data" / "test_data.json")))
    if not path.is_file():
        raise ValueError(f"Test data file does not exist: {path}")
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows)
            values = next(rows)
            data = dict(zip(headers, values))
        except StopIteration as exc:
            raise ValueError("Excel file needs a header row and one data row") from exc
        finally:
            workbook.close()
    else:
        raise ValueError("TEST_DATA must point to a .json or .xlsx file")

    if not isinstance(data, dict):
        raise ValueError("Test data must contain a JSON object or an Excel data row")
    missing = [key for key in FIELDS if data.get(key) is None or str(data[key]).strip() == ""]
    if missing:
        raise ValueError(f"Missing test fields: {', '.join(missing)}")
    data["quantity"] = int(data["quantity"])
    if data["quantity"] < 2:
        raise ValueError("quantity must be at least 2 to test a cart update")
    for field in ("base_url", "product", "first_name", "last_name", "telephone"):
        data[field] = str(data[field]).strip()
    url = urlparse(data["base_url"])
    if url.scheme != "https" or url.netloc != "tutorialsninja.com":
        raise ValueError("This project's locators require https://tutorialsninja.com/demo/")
    return data
